"""
Regression tests for the bank-to-GL reconciliation engine.

Three layers, each catching a different kind of breakage:

  1. Golden numbers   — the full June 2026 result is pinned. Any change to the
                        output at all fails here.
  2. Named edge cases — one test per planted scenario, so when a golden number
                        breaks, the test name says WHICH rule broke.
  3. Controls         — prove the proof is not circular, the generator is
                        deterministic, and the files are valid UTF-8.

The tests call reconcile(), not the CSV reader directly. When the data layer
moves to SQLite (Phase 2), these tests must still pass unchanged.

Run from the repo root:
    python -m pytest -v
"""

import csv
import filecmp
import subprocess
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "src"))

import reconcile as rec  # noqa: E402

DATA = ROOT / "data"
BANK = DATA / "bank_statement_jun2026.csv"
GL = DATA / "gl_cash_extract_jun2026.csv"
BALANCES = DATA / "balances_jun2026.csv"
DATA_FILES = [BANK.name, GL.name, BALANCES.name]


# ------------------------------------------------------------------ fixtures + helpers
@pytest.fixture(scope="module")
def result():
    bank, gl, balances, matches, exceptions, proof = rec.reconcile(BANK, GL, BALANCES)
    return {"bank": bank, "gl": gl, "balances": balances,
            "matches": matches, "exceptions": exceptions, "proof": proof}


def match_for_gl_memo(result, text):
    """The single match whose GL memo contains `text`."""
    gl = result["gl"]
    found = [m for m in result["matches"] if text in gl.at[m["gl_idx"], "Memo"]]
    assert len(found) == 1, f"expected one match for {text!r}, got {len(found)}"
    return found[0]


def exception_for(result, text):
    """The single exception whose description contains `text`."""
    found = [e for e in result["exceptions"] if text in str(e["Description"])]
    assert len(found) == 1, f"expected one exception for {text!r}, got {len(found)}"
    return found[0]


def iso(ts):
    return ts.date().isoformat()


def write_rows(path, rows):
    with open(path, "w", newline="", encoding="utf-8") as f:
        csv.writer(f).writerows(rows)


def read_rows(path):
    with open(path, newline="", encoding="utf-8") as f:
        return list(csv.reader(f))


# ------------------------------------------------------------------ 1. golden numbers
def test_row_counts(result):
    assert len(result["bank"]) == 127
    assert len(result["gl"]) == 127


def test_match_counts_per_pass(result):
    by_method = {}
    for m in result["matches"]:
        key = (m["pass"], m["method"])
        by_method[key] = by_method.get(key, 0) + 1
    assert by_method == {
        ("Exact", "amount + date"): 95,
        ("Timing", "check no."): 13,
        ("Timing", "date window"): 2,
        ("Tolerance", "fuzzy"): 8,
    }
    assert len(result["matches"]) == 118


def test_matching_is_one_to_one(result):
    banks = [m["bank_idx"] for m in result["matches"]]
    gls = [m["gl_idx"] for m in result["matches"]]
    assert len(banks) == len(set(banks))
    assert len(gls) == len(set(gls))


def test_every_row_is_matched_or_an_exception(result):
    n_bank_exc = sum(1 for e in result["exceptions"] if e["Side"] == "Bank only")
    n_gl_exc = sum(1 for e in result["exceptions"] if e["Side"] == "GL only")
    assert len(result["matches"]) + n_bank_exc == len(result["bank"])
    assert len(result["matches"]) + n_gl_exc == len(result["gl"])


# Every exception, in rank order: (side, amount in cents, category).
# A wrong diagnosis fails here, not just a wrong count.
GOLDEN_EXCEPTIONS = [
    ("GL only",    4831077, "DIT"),
    ("GL only",   -1589023, "OS_CHECK"),
    ("GL only",    1276431, "DIT"),
    ("GL only",   -1023018, "DUPLICATE"),
    ("GL only",     940658, "DEP_NOT_CREDITED"),
    ("GL only",    -741266, "OS_CHECK"),
    ("GL only",    -417809, "OS_CHECK"),
    ("Bank only",  -387540, "NSF_RETURN"),
    ("GL only",    -234055, "OS_CHECK"),
    ("Bank only",   215000, "UNID_CREDIT"),
    ("GL only",    -128643, "PMT_IN_TRANSIT"),
    ("Bank only",  -126000, "UNID_DEBIT"),
    ("Bank only",   -61208, "INTEREST_EXP"),
    ("Bank only",   -38944, "BANK_CHARGE"),
    ("Bank only",    21487, "INTEREST_INC"),
    ("Bank only",   -12500, "BANK_CHARGE"),
    ("Bank only",    -4800, "BANK_CHARGE"),
    ("Bank only",    -4500, "BANK_CHARGE"),
    ("Residual",       -35, "RESIDUAL"),      # cents pass 3 could not pair away
]


def test_exception_list_exact(result):
    actual = [(e["Side"], e["cents"], e["Category"]) for e in result["exceptions"]]
    assert actual == GOLDEN_EXCEPTIONS


def test_ranks_are_by_dollar_exposure(result):
    exposures = [abs(e["cents"]) for e in result["exceptions"]]
    assert exposures == sorted(exposures, reverse=True)
    assert [e["Rank"] for e in result["exceptions"]] == list(range(1, 20))


def test_proof_ties_to_zero(result):
    proof = result["proof"]
    assert proof["completeness"]["Bank"]["diff"] == 0
    assert proof["completeness"]["GL"]["diff"] == 0
    assert proof["difference"] == 0
    assert proof["adj_bank"] == proof["adj_gl"] == 43_396_047   # $433,960.47


def test_pass3_residual_is_carried_to_proof(result):
    """The residual is a reconciling item in its own right: it stays on the rec
    until the write-off JE is booked, which is how the NEXT period inherits it."""
    residual = dict(result["proof"]["gl_lines"])["add: pass-3 amount residuals (pending write-off JE)"]
    assert residual == -35                                       # −$0.35, not lost
    items = [e for e in result["exceptions"] if e["Category"] == "RESIDUAL"]
    assert len(items) == 1 and items[0]["cents"] == -35


# ------------------------------------------------------------------ 2. named edge cases
def test_late_check_matched_by_number(result):
    """Clears 11 days after booking: outside ±5 days, only the check number pairs it."""
    m = match_for_gl_memo(result, "Telecon Design — survey invoice")
    assert m["method"] == "check no."
    assert m["date_delta"] == 11


def test_competing_leases_assigned_optimally(result):
    """Greedy nearest-date would pair Jun 10 with Jun 11 and strand Jun 15."""
    bank = result["bank"]
    pairs = {}
    for unit in (1, 2):
        m = match_for_gl_memo(result, f"vehicle lease unit {unit}")
        assert m["method"] == "date window"
        pairs[unit] = iso(bank.at[m["bank_idx"], "Date"])
    assert pairs == {1: "2026-06-10", 2: "2026-06-15"}


def test_recurring_payment_not_flagged_duplicate(result):
    """Same vendor, same amount, different document: a recurring payment, not an error."""
    e = exception_for(result, "vehicle lease unit 3")
    assert e["Category"] == "PMT_IN_TRANSIT"


def test_true_duplicate_still_caught(result):
    """Same document, same amount as a matched entry: reverse it."""
    dups = [e for e in result["exceptions"] if e["Category"] == "DUPLICATE"]
    assert len(dups) == 1
    assert dups[0]["Doc/Ref"] == "CHQ#1042"
    assert dups[0]["cents"] == -1_023_018


def test_interest_charged_is_expense(result):
    """A debit containing INTEREST must never be booked as interest income."""
    assert exception_for(result, "INTEREST ON OPERATING LINE OF CREDIT")["Category"] == "INTEREST_EXP"
    assert exception_for(result, "INTEREST EARNED")["Category"] == "INTEREST_INC"


def test_nsf_return_vs_nsf_fee(result):
    """The fee is a bank charge; the returned cheque reverses a receivable."""
    assert exception_for(result, "NSF RETURNED ITEM FEE")["Category"] == "BANK_CHARGE"
    assert exception_for(result, "RETURNED ITEM NSF CUSTOMER")["Category"] == "NSF_RETURN"


def test_midmonth_deposit_not_in_transit(result):
    """Deposits clear in a day or two; one from Jun 12 still missing is a problem."""
    assert exception_for(result, "change order")["Category"] == "DEP_NOT_CREDITED"
    for text in ("Rogers Communications — progress billing", "Cogeco underground build"):
        dits = [e for e in result["exceptions"]
                if e["Description"] == text and e["Category"] == "DIT"]
        assert len(dits) == 1


def test_fuzzy_pass_never_crosses_direction(result):
    bank, gl = result["bank"], result["gl"]
    for m in result["matches"]:
        assert (bank.at[m["bank_idx"], "cents"] > 0) == (gl.at[m["gl_idx"], "cents"] > 0)


# ------------------------------------------------------------------ 3. controls
def test_missing_bank_row_breaks_proof(tmp_path):
    """The proof starts from stated balances, so a lost row cannot hide."""
    rows = read_rows(BANK)
    removed = rows.pop(40)
    bank_missing = tmp_path / "bank.csv"
    write_rows(bank_missing, rows)

    *_, proof = rec.reconcile(bank_missing, GL, BALANCES)
    # diff = stated closing − computed closing: exactly the amount of the lost row
    missing = rec.to_cents(removed[3])
    assert proof["completeness"]["Bank"]["diff"] == missing
    assert proof["completeness"]["GL"]["diff"] == 0
    assert proof["difference"] == missing


def test_missing_gl_row_breaks_proof(tmp_path):
    rows = read_rows(GL)
    removed = rows.pop(40)
    gl_missing = tmp_path / "gl.csv"
    write_rows(gl_missing, rows)

    *_, proof = rec.reconcile(BANK, gl_missing, BALANCES)
    assert proof["completeness"]["GL"]["diff"] == rec.to_cents(removed[4])
    assert proof["difference"] != 0


def test_wrong_stated_balance_breaks_proof(tmp_path):
    rows = read_rows(BALANCES)
    header = rows[0]
    for row in rows[1:]:
        if row[0] == "BANK":
            i = header.index("ClosingBalance")
            row[i] = f"{rec.to_cents(row[i]) / 100 + 0.01:.2f}"   # off by one cent
    tampered = tmp_path / "balances.csv"
    write_rows(tampered, rows)

    *_, proof = rec.reconcile(BANK, GL, tampered)
    assert proof["completeness"]["Bank"]["diff"] == 1
    assert proof["difference"] == 1


def test_every_category_has_a_proof_line(monkeypatch, result):
    """A new category without a proof line must fail loudly, not vanish."""
    monkeypatch.setitem(rec.CATEGORIES, "NEW_UNPLACED", ("cause", "action"))
    with pytest.raises(RuntimeError):
        rec.build_proof(result["bank"], result["gl"], result["matches"],
                        result["exceptions"], result["balances"])


def test_to_cents_is_exact():
    assert rec.to_cents("0.29") == 29           # int(0.29 * 100) would give 28
    assert rec.to_cents("-1286.43") == -128643
    assert rec.to_cents("1,234.56") == 123456


def _generate(out_dir):
    subprocess.run([sys.executable, str(ROOT / "src" / "generate_data.py"),
                    "--out-dir", str(out_dir)],
                   check=True, capture_output=True)


def test_generator_is_deterministic(tmp_path):
    run_a, run_b = tmp_path / "a", tmp_path / "b"
    _generate(run_a)
    _generate(run_b)
    for name in DATA_FILES:
        assert filecmp.cmp(run_a / name, run_b / name, shallow=False), name


def test_committed_data_matches_generator(tmp_path):
    """The CSVs in data/ are exactly what the generator produces."""
    _generate(tmp_path)
    for name in DATA_FILES:
        assert filecmp.cmp(tmp_path / name, DATA / name, shallow=False), \
            f"{name} differs from the generator output — regenerate data/"


@pytest.mark.parametrize("name", DATA_FILES)
def test_data_files_are_utf8(name):
    (DATA / name).read_bytes().decode("utf-8")      # raises on cp1252 bytes like 0x97


def test_report_builds(result, tmp_path):
    out = rec.build_report(result["bank"], result["gl"], result["matches"],
                           result["exceptions"], result["proof"],
                           result["balances"]["period_end"], tmp_path / "rec.xlsx")
    wb = load_workbook(out)
    assert wb.sheetnames == ["Summary", "Exceptions", "Matched — Exact", "Matched — Timing",
                             "Matched — Tolerance", "Bank Statement (source)", "GL Extract (source)"]
    summary_text = [c.value for row in wb["Summary"].iter_rows() for c in row if c.value]
    assert "Status: TIES" in summary_text
    assert wb["Exceptions"].max_row == 4 + 19      # header row 4 + 19 exceptions