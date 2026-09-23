"""
reconcile.py — Bank-to-GL three-pass reconciliation engine.

Reads a bank statement CSV, a GL cash-account extract CSV and a balances CSV
(stated opening/closing balances from the statement header and the trial
balance), matches the transactions in three passes of decreasing strictness,
classifies whatever is left as exceptions with a probable cause and a
suggested action, and writes a multi-tab, close-ready Excel workbook.

    Pass 1 — EXACT      same signed amount, same date
    Pass 2 — TIMING     same signed amount, cleared later:
                          a) checks: same check number, any lag within the period
                          b) others: dates within ±5 days, assigned optimally
                             when several same-amount rows compete
    Pass 3 — TOLERANCE  amounts within $0.99, dates within ±7 days,
                        fuzzy description similarity (rounding / keying errors)

All money is handled as integer cents, parsed with Decimal straight from the
CSV text; floats appear only when values are written to the report.

The reconciliation proof starts from the STATED balances, not from sums of the
rows, and two completeness controls check that each file's rows actually roll
the stated opening balance to the stated closing balance. A missing or extra
row therefore breaks the proof instead of disappearing inside it.

Usage:
    python src/reconcile.py \
        --bank     data/bank_statement_jun2026.csv \
        --gl       data/gl_cash_extract_jun2026.csv \
        --balances data/balances_jun2026.csv \
        --out      output/reconciliation_report_jun2026.xlsx
"""

import argparse
import csv
import re
from collections import defaultdict
from datetime import date
from decimal import ROUND_HALF_UP, Decimal
from difflib import SequenceMatcher
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ------------------------------------------------------------------ parameters
TIMING_WINDOW_DAYS = 5      # pass 2b: max days between GL booking and bank clearing
TOLERANCE_DOLLARS = 0.99    # pass 3: max absolute amount difference
TOLERANCE_WINDOW_DAYS = 7   # pass 3: max days apart
FUZZY_THRESHOLD = 0.35      # pass 3: min description similarity (0..1)
MONTH_END_DAYS = 5          # "near month end" = within N days of the period end
ENTITY_LABEL = "Operating Account 1010 — Cash"

# Check references on either side, e.g. "CHQ#1083". Adapt to the bank / ERP format.
CHECK_REF = re.compile(r"^\s*CHQ\s*#?\s*(\d+)\s*$", re.I)

STOPWORDS = {"EFT", "PAD", "POS", "DEP", "CHQ", "ADP", "PPD", "INC", "LTD",
             "SVC", "PMT", "PAYMENT", "INVOICE", "THE", "OF", "AND", "CDA",
             "CANADA", "REF"}


# ------------------------------------------------------------------ money + text helpers
def to_cents(text) -> int:
    """'-1286.43' -> -128643, exactly. Never goes through a float."""
    d = Decimal(str(text).strip().replace(",", ""))
    return int((d * 100).quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def dollars(c: int) -> float:
    """Display only: cents -> dollars for the Excel report."""
    return c / 100


def check_number(ref) -> str:
    """'CHQ#1083' -> '1083'; anything else -> '' (not None: pandas turns None into NaN)."""
    m = CHECK_REF.match(str(ref)) if pd.notna(ref) else None
    return m.group(1) if m else ""


def normalize(desc: str) -> str:
    s = re.sub(r"[^A-Z0-9 ]", " ", str(desc).upper())
    s = re.sub(r"\b(REF|CHQ)?\d+\b", " ", s)          # strip reference numbers
    tokens = [t for t in s.split() if t not in STOPWORDS and len(t) > 2]
    return " ".join(tokens)


def similarity(a: str, b: str) -> float:
    """Blend of sequence ratio and token overlap — robust to reordering."""
    na, nb = normalize(a), normalize(b)
    if not na or not nb:
        return 0.0
    seq = SequenceMatcher(None, na, nb).ratio()
    ta, tb = set(na.split()), set(nb.split())
    jac = len(ta & tb) / len(ta | tb) if ta | tb else 0.0
    return max(seq, jac)


# ------------------------------------------------------------------ loading
def load(bank_path: str, gl_path: str):
    bank = pd.read_csv(bank_path, dtype={"Amount": str, "Reference": str},
                       parse_dates=["Date"], encoding="utf-8-sig")
    gl = pd.read_csv(gl_path, dtype={"Amount": str, "DocNo": str},
                     parse_dates=["Date"], encoding="utf-8-sig")
    for df, ref_col in ((bank, "Reference"), (gl, "DocNo")):
        df["cents"] = df["Amount"].map(to_cents)
        df["Amount"] = df["cents"].map(dollars)            # display copy
        df["chk"] = df[ref_col].map(check_number)
        df["src_row"] = df.index + 2                       # row in the source CSV (after header)
    return bank.reset_index(drop=True), gl.reset_index(drop=True)


def load_balances(path: str) -> dict:
    """Stated balances — the independent figures the proof is built on."""
    out = {}
    with open(path, newline="", encoding="utf-8-sig") as f:
        for r in csv.DictReader(f):
            out[r["Source"].strip().upper()] = {
                "start": date.fromisoformat(r["PeriodStart"]),
                "end": date.fromisoformat(r["PeriodEnd"]),
                "open": to_cents(r["OpeningBalance"]),
                "close": to_cents(r["ClosingBalance"]),
            }
    missing = {"BANK", "GL"} - out.keys()
    if missing:
        raise ValueError(f"balances file missing: {sorted(missing)}")
    if out["BANK"]["end"] != out["GL"]["end"]:
        raise ValueError("bank and GL balances are for different period ends")
    out["period_end"] = pd.Timestamp(out["BANK"]["end"])
    return out


# ------------------------------------------------------------------ matching
def optimal_date_pairs(bank, gl, bank_ids, gl_ids, window, compatible):
    """
    Pair same-amount rows by date, one-to-one, maximising the number of pairs
    and then minimising the total day gap.

    Greedy nearest-date fails when rows compete: a Jun 10 debit grabs the
    Jun 11 booking (gap 1) and strands a Jun 15 debit 7 days from the Jun 8
    booking. The optimum pairs Jun 8→10 and Jun 11→15. On a line, an optimal
    matching never needs crossing pairs, so a DP over both lists sorted by
    date finds it exactly. Groups are tiny (same exact amount), so this is cheap.
    """
    bs = sorted(bank_ids, key=lambda i: (bank.at[i, "Date"], i))
    gs = sorted(gl_ids, key=lambda i: (gl.at[i, "Date"], i))
    n, m = len(bs), len(gs)
    best = [[(0, 0)] * (m + 1) for _ in range(n + 1)]   # (pairs, -total_gap)
    move = [[None] * (m + 1) for _ in range(n + 1)]
    for i in range(n - 1, -1, -1):
        for j in range(m - 1, -1, -1):
            options = []
            gap = abs((bank.at[bs[i], "Date"] - gl.at[gs[j], "Date"]).days)
            if gap <= window and compatible(bs[i], gs[j]):
                k, neg = best[i + 1][j + 1]
                options.append(((k + 1, neg - gap), "pair"))
            options.append((best[i + 1][j], "skip_bank"))
            options.append((best[i][j + 1], "skip_gl"))
            best[i][j], move[i][j] = max(options, key=lambda o: o[0])
    pairs, i, j = [], 0, 0
    while i < n and j < m:
        if move[i][j] == "pair":
            pairs.append((bs[i], gs[j]))
            i, j = i + 1, j + 1
        elif move[i][j] == "skip_bank":
            i += 1
        else:
            j += 1
    return pairs


def run_matching(bank: pd.DataFrame, gl: pd.DataFrame):
    """One-to-one matching in three passes. Returns (matches, bank_open, gl_open)."""
    matches = []
    bank_open = set(bank.index)
    gl_open = set(gl.index)

    def compatible(bi, gi):
        """Two rows that both carry a check number must carry the same one."""
        bc, gc = bank.at[bi, "chk"], gl.at[gi, "chk"]
        return not bc or not gc or bc == gc

    def pair(bi, gi, pass_name, method, score=None):
        b, g = bank.loc[bi], gl.loc[gi]
        matches.append({
            "pass": pass_name, "method": method,
            "bank_idx": bi, "gl_idx": gi,
            "date_delta": int((b["Date"] - g["Date"]).days),
            "cents_delta": int(b["cents"] - g["cents"]),
            "score": score,
        })
        bank_open.discard(bi)
        gl_open.discard(gi)

    # Pass 1 — exact: same cents, same date
    gl_exact = defaultdict(list)
    for gi in sorted(gl_open):
        gl_exact[(gl.at[gi, "cents"], gl.at[gi, "Date"])].append(gi)
    for bi in sorted(bank_open):
        for gi in gl_exact.get((bank.at[bi, "cents"], bank.at[bi, "Date"]), []):
            if gi in gl_open and compatible(bi, gi):
                pair(bi, gi, "Exact", "amount + date")
                break

    # Pass 2a — checks: same check number and same cents, whatever the clearing lag
    gl_checks = defaultdict(list)
    for gi in sorted(gl_open):
        if gl.at[gi, "chk"]:
            gl_checks[(gl.at[gi, "chk"], gl.at[gi, "cents"])].append(gi)
    for bi in sorted(bank_open):
        if bank.at[bi, "chk"]:
            for gi in gl_checks.get((bank.at[bi, "chk"], bank.at[bi, "cents"]), []):
                if gi in gl_open:
                    pair(bi, gi, "Timing", "check no.")
                    break

    # Pass 2b — same cents within the window, optimal assignment per amount
    by_amount = defaultdict(lambda: ([], []))
    for bi in sorted(bank_open):
        by_amount[bank.at[bi, "cents"]][0].append(bi)
    for gi in sorted(gl_open):
        if gl.at[gi, "cents"] in by_amount:
            by_amount[gl.at[gi, "cents"]][1].append(gi)
    for cents_key in sorted(by_amount):
        bank_ids, gl_ids = by_amount[cents_key]
        if bank_ids and gl_ids:
            for bi, gi in optimal_date_pairs(bank, gl, bank_ids, gl_ids,
                                             TIMING_WINDOW_DAYS, compatible):
                pair(bi, gi, "Timing", "date window")

    # Pass 3 — tolerance + fuzzy description: best candidates first, globally
    tol_cents = round(TOLERANCE_DOLLARS * 100)
    candidates = []
    for bi in sorted(bank_open):
        b = bank.loc[bi]
        for gi in sorted(gl_open):
            g = gl.loc[gi]
            if (b["cents"] > 0) != (g["cents"] > 0):
                continue  # never pair money in with money out
            if abs(b["cents"] - g["cents"]) > tol_cents:
                continue
            days = abs((b["Date"] - g["Date"]).days)
            if days > TOLERANCE_WINDOW_DAYS or not compatible(bi, gi):
                continue
            score = similarity(b["Description"], g["Memo"])
            if score >= FUZZY_THRESHOLD:
                candidates.append((-score, days, abs(b["cents"] - g["cents"]), bi, gi, score))
    for _, _, _, bi, gi, score in sorted(candidates):
        if bi in bank_open and gi in gl_open:
            pair(bi, gi, "Tolerance", "fuzzy", round(score, 2))

    return matches, bank_open, gl_open


# ------------------------------------------------------------------ exceptions
# Machine-readable category codes drive the proof and the report styling;
# the text is for people. Each code: (probable cause, suggested action).
CATEGORIES = {
    # bank only — need a journal entry
    "BANK_CHARGE": ("Bank charge — not booked in GL",
                    "Book JE: DR 6220 Bank Charges / CR 1010 Cash"),
    "INTEREST_INC": ("Interest earned — not booked in GL",
                     "Book JE: DR 1010 Cash / CR 4210 Interest Income"),
    "INTEREST_EXP": ("Interest charged — not booked in GL",
                     "Book JE: DR 7110 Interest Expense / CR 1010 Cash"),
    "NSF_RETURN": ("Customer deposit returned NSF — not booked in GL",
                   "Book JE: DR 1200 Accounts Receivable / CR 1010 Cash — contact customer"),
    # bank only — need investigation
    "UNID_CREDIT": ("Unidentified bank credit",
                    "Trace with bank / AR — identify payer before booking"),
    "UNID_DEBIT": ("Unidentified bank debit",
                   "Investigate with bank — possible unauthorized PAD"),
    # GL only — timing, carried on the rec
    "OS_CHECK": ("Outstanding check",
                 "Carry as reconciling item — follow up if stale > 60 days"),
    "PMT_IN_TRANSIT": ("Payment initiated near month end, not yet cleared",
                       "Confirm settlement on next bank statement"),
    "DIT": ("Deposit in transit",
            "Verify credit on next bank statement"),
    # GL only — need investigation
    "PMT_NOT_DEBITED": ("Payment booked mid-period, never debited by bank",
                        "Investigate — confirm the payment was actually sent"),
    "DEP_NOT_CREDITED": ("Deposit booked mid-period, never credited by bank",
                         "Investigate — confirm funds were actually deposited"),
    "DUPLICATE": ("Possible duplicate posting (same document and amount as another entry)",
                  "Review source document — reverse the duplicate JE"),
    # not a row in either file: the cents left over by pass 3, which stay a
    # reconciling item until somebody books the write-off
    "RESIDUAL": ("Pass-3 amount residuals — pending write-off JE",
                 "Book JE: write off the rounding difference to 6900 Misc Expense"),
}

INTEREST_PATTERN = re.compile(r"\bINTEREST\b", re.I)
FEE_PATTERN = re.compile(r"\bFEE\b|SERVICE CHARGE|SVC CHG|OVERDRAFT", re.I)
NSF_PATTERN = re.compile(r"\bNSF\b|RETURNED (ITEM|CHQ|CHEQUE|CHECK)", re.I)


def classify_bank_only(b) -> str:
    desc = str(b["Description"])
    if INTEREST_PATTERN.search(desc):
        return "INTEREST_INC" if b["cents"] > 0 else "INTEREST_EXP"
    if FEE_PATTERN.search(desc):          # checked before NSF: "NSF RETURNED ITEM FEE" is a fee
        return "BANK_CHARGE"
    if NSF_PATTERN.search(desc) and b["cents"] < 0:
        return "NSF_RETURN"
    return "UNID_CREDIT" if b["cents"] > 0 else "UNID_DEBIT"


def residual_item(residual_cents, period_end):
    """The cents pass 3 could not pair away, as an exception row of its own.

    Carrying it as an item rather than a footnote is what keeps the NEXT
    period honest: until the write-off JE is booked, the difference is still
    there, and the proof should say so."""
    cause, action = CATEGORIES["RESIDUAL"]
    return {"Side": "Residual", "Date": period_end,
            "Description": "Pass-3 amount residuals across matched pairs",
            "Doc/Ref": "", "cents": int(residual_cents), "Amount": dollars(int(residual_cents)),
            "Category": "RESIDUAL", "Probable cause": cause, "Suggested action": action,
            "Source row": 0}


def classify_exceptions(bank, gl, bank_open, gl_open, matches, period_end):
    def doc_key(gi):
        doc = gl.at[gi, "DocNo"]
        doc = "" if pd.isna(doc) else str(doc).strip().upper()
        return (int(gl.at[gi, "cents"]), doc) if doc else None

    # A duplicate repeats the SAME document for the same amount. Recurring
    # payments (same vendor, same amount) carry different documents.
    seen_docs = {doc_key(m["gl_idx"]) for m in matches} - {None}
    rows = []

    for bi in sorted(bank_open):
        b = bank.loc[bi]
        rows.append({"Side": "Bank only", "Date": b["Date"], "Description": b["Description"],
                     "Doc/Ref": b["Reference"], "cents": int(b["cents"]),
                     "Category": classify_bank_only(b), "Source row": int(b["src_row"])})

    for gi in sorted(gl_open, key=lambda i: (gl.at[i, "Date"], i)):
        g = gl.loc[gi]
        key = doc_key(gi)
        near_eom = (period_end - g["Date"]).days <= MONTH_END_DAYS
        if key is not None and key in seen_docs:
            cat = "DUPLICATE"
        elif g["cents"] < 0 and g["chk"]:
            cat = "OS_CHECK"
        elif g["cents"] < 0:
            cat = "PMT_IN_TRANSIT" if near_eom else "PMT_NOT_DEBITED"
        else:
            cat = "DIT" if near_eom else "DEP_NOT_CREDITED"
        if key is not None:
            seen_docs.add(key)
        rows.append({"Side": "GL only", "Date": g["Date"], "Description": g["Memo"],
                     "Doc/Ref": g["DocNo"], "cents": int(g["cents"]),
                     "Category": cat, "Source row": int(g["src_row"])})

    for r in rows:
        r["Amount"] = dollars(r["cents"])
        cause, action = CATEGORIES[r["Category"]]
        if r["Category"] == "OS_CHECK" and (period_end - r["Date"]).days <= MONTH_END_DAYS:
            cause += " (issued near month end)"
        r["Probable cause"], r["Suggested action"] = cause, action

    rows.sort(key=lambda r: (-abs(r["cents"]), r["Date"], r["Side"], r["Source row"]))
    for rank, r in enumerate(rows, 1):
        r["Rank"] = rank
    return rows


def merge_carried(exceptions, still_open, period_end):
    """Fold items carried from earlier periods into this period's exception list.

    A bank rec shows EVERY outstanding item, not only the ones raised this
    month: a check written in June is still a reconciling item in July. They
    are ranked together by exposure and re-numbered."""
    rows = list(exceptions)
    for item in still_open:
        cause, action = CATEGORIES[item["category"]]
        date_ = pd.Timestamp(item["origin_date"]) if item["origin_date"] else period_end
        side = {"BANK": "Bank only", "GL": "GL only"}.get(item["side"], "Residual")
        rows.append({
            "Side": side,
            "Date": date_,
            "Description": item["description"] or CATEGORIES[item["category"]][0],
            "Doc/Ref": item["doc_ref"] or "",
            "cents": int(item["amount_cents"]), "Amount": dollars(int(item["amount_cents"])),
            "Category": item["category"], "Probable cause": cause + " (carried forward)",
            "Suggested action": action, "Source row": int(item["source_row"] or 0),
            "Origin period": item["origin_period_id"], "item_id": item["item_id"],
        })
    for r in rows:
        r.setdefault("Origin period", str(period_end)[:7])
        r["Age (days)"] = int((period_end - r["Date"]).days)
    rows.sort(key=lambda r: (-abs(r["cents"]), r["Date"], r["Side"], r["Source row"]))
    for rank, r in enumerate(rows, 1):
        r["Rank"] = rank
    return rows


# ------------------------------------------------------------------ proof
# Where each category sits on the rec. Every category must appear exactly once;
# build_proof refuses to run otherwise, so a new category can't silently vanish.
BANK_SIDE = [
    ("add: deposits in transit", ["DIT"]),
    ("add: deposits not credited by bank (investigate)", ["DEP_NOT_CREDITED"]),
    ("less: outstanding checks", ["OS_CHECK"]),
    ("less: payments in transit", ["PMT_IN_TRANSIT"]),
    ("less: payments not debited by bank (investigate)", ["PMT_NOT_DEBITED"]),
]
GL_SIDE = [
    ("add: pass-3 amount residuals (pending write-off JE)", ["RESIDUAL"]),
    ("add: bank charges not booked", ["BANK_CHARGE"]),
    ("add: interest not booked (net)", ["INTEREST_INC", "INTEREST_EXP"]),
    ("add: NSF returned deposits not booked", ["NSF_RETURN"]),
    ("add: unidentified bank items (pending ID)", ["UNID_CREDIT", "UNID_DEBIT"]),
]
REVERSED_ON_GL = [("add back: duplicate postings to reverse", ["DUPLICATE"])]


def build_proof(bank, gl, matches, exceptions, balances) -> dict:
    placed = [c for _, cats in BANK_SIDE + GL_SIDE + REVERSED_ON_GL for c in cats]
    if sorted(placed) != sorted(CATEGORIES):
        raise RuntimeError("every exception category must sit on exactly one proof line")

    by_cat = defaultdict(int)
    for e in exceptions:
        by_cat[e["Category"]] += e["cents"]

    def line_total(cats):
        return sum(by_cat[c] for c in cats)

    B, G = balances["BANK"], balances["GL"]
    completeness = {}
    for side, df, bal in (("Bank", bank, B), ("GL", gl, G)):
        movement = int(df["cents"].sum())
        completeness[side] = {"open": bal["open"], "movement": movement,
                              "computed_close": bal["open"] + movement,
                              "stated_close": bal["close"],
                              "diff": bal["close"] - (bal["open"] + movement)}

    bank_lines = [(label, line_total(cats)) for label, cats in BANK_SIDE]
    gl_lines = [(label, line_total(cats)) for label, cats in GL_SIDE]
    gl_lines += [(label, -line_total(cats)) for label, cats in REVERSED_ON_GL]

    adj_bank = B["close"] + sum(v for _, v in bank_lines)
    adj_gl = G["close"] + sum(v for _, v in gl_lines)
    return {"completeness": completeness,
            "bank_close": B["close"], "bank_lines": bank_lines, "adj_bank": adj_bank,
            "gl_close": G["close"], "gl_lines": gl_lines, "adj_gl": adj_gl,
            "difference": adj_bank - adj_gl}


# ------------------------------------------------------------------ excel report
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)
TITLE_FONT = Font(bold=True, size=14, color="1F3864")
SUB_FONT = Font(size=10, color="595959")
MONEY = '#,##0.00;[Red](#,##0.00)'
COUNT = "0"
PCT = "0.0%"

_JE, _INVESTIGATE = "E2EFDA", "FFC7CE"
FLAG_FILLS = {"ESCALATE": PatternFill("solid", fgColor="FFC7CE"),
              "STALE": PatternFill("solid", fgColor="F8CBAD"),
              "FOLLOW_UP": PatternFill("solid", fgColor="FFF2CC")}

CATEGORY_FILLS = {code: PatternFill("solid", fgColor=color) for code, color in {
    "DUPLICATE": "F8CBAD",
    "UNID_CREDIT": _INVESTIGATE, "UNID_DEBIT": _INVESTIGATE,
    "PMT_NOT_DEBITED": _INVESTIGATE, "DEP_NOT_CREDITED": _INVESTIGATE,
    "OS_CHECK": "FFF2CC", "PMT_IN_TRANSIT": "FFF2CC",
    "DIT": "DDEBF7",
    "BANK_CHARGE": _JE, "INTEREST_INC": _JE, "INTEREST_EXP": _JE, "NSF_RETURN": _JE,
}.items()}


def style_header(ws, row, ncols, freeze=True):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = BORDER
        cell.alignment = Alignment(vertical="center")
    if freeze:
        ws.freeze_panes = ws.cell(row=row + 1, column=1)


def autosize(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def write_table(ws, start_row, headers, rows, money_cols=(), date_cols=(), freeze=True):
    for c, h in enumerate(headers, 1):
        ws.cell(row=start_row, column=c, value=h)
    style_header(ws, start_row, len(headers), freeze)
    r = start_row
    for rowdata in rows:
        r += 1
        for c, v in enumerate(rowdata, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border = BORDER
            if c in money_cols:
                cell.number_format = MONEY
            if c in date_cols:
                cell.number_format = "yyyy-mm-dd"
    return r


def write_label_value_table(ws, start_row, header, rows, bold_prefixes=()):
    """rows: (label, value, number_format). Formats are per row, not per column."""
    r = write_table(ws, start_row, list(header), [(lbl, val) for lbl, val, _ in rows], freeze=False)
    for i, (label, _, fmt) in enumerate(rows, start_row + 1):
        ws.cell(row=i, column=2).number_format = fmt
        if str(label).startswith(bold_prefixes):
            ws.cell(row=i, column=1).font = Font(bold=True)
            ws.cell(row=i, column=2).font = Font(bold=True)
    return r


def build_report(bank, gl, matches, exceptions, proof, period_end, out_path, ledger=None):
    wb = Workbook()
    period_label = period_end.strftime("%B %Y")

    by_pass = defaultdict(list)
    for m in matches:
        by_pass[m["pass"]].append(m)
    n_check = sum(1 for m in by_pass["Timing"] if m["method"] == "check no.")

    # ---------------- Summary tab
    ws = wb.active
    ws.title = "Summary"
    ws.sheet_properties.tabColor = "1F3864"
    ws["A1"] = "Bank-to-GL Reconciliation"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"{ENTITY_LABEL}  ·  Period: {period_label}"
    ws["A2"].font = SUB_FONT
    ws["A3"] = "Engine: three-pass matcher (exact / timing / tolerance+fuzzy)"
    ws["A3"].font = SUB_FONT

    stats = [
        ("Bank transactions", len(bank), COUNT),
        ("GL transactions", len(gl), COUNT),
        ("Matched — Pass 1 (exact)", len(by_pass["Exact"]), COUNT),
        ("Matched — Pass 2 (timing)", len(by_pass["Timing"]), COUNT),
        ("   of which by check number", n_check, COUNT),
        (f"   of which by date window ±{TIMING_WINDOW_DAYS} days", len(by_pass["Timing"]) - n_check, COUNT),
        ("Matched — Pass 3 (tolerance ≤ $0.99 + fuzzy)", len(by_pass["Tolerance"]), COUNT),
        ("Total reconciled", len(matches), COUNT),
        ("Match rate (bank side)", len(matches) / len(bank) if len(bank) else 0, PCT),
        ("Exceptions isolated", len(exceptions), COUNT),
        ("Exception exposure (gross $)", dollars(sum(abs(e["cents"]) for e in exceptions)), MONEY),
    ]
    r = write_label_value_table(ws, 5, ("Metric", "Value"), stats,
                                bold_prefixes=("Total reconciled",))

    controls = []
    for side, c in proof["completeness"].items():
        label = "bank statement" if side == "Bank" else "general ledger"
        controls += [
            (f"{side}: opening balance per {label}", dollars(c["open"]), MONEY),
            (f"{side}: + net movement of rows in file", dollars(c["movement"]), MONEY),
            (f"{side}: = computed closing balance", dollars(c["computed_close"]), MONEY),
            (f"{side}: closing balance per {label}", dollars(c["stated_close"]), MONEY),
            (f"{side}: difference (missing / extra rows)", dollars(c["diff"]), MONEY),
        ]
    r = write_label_value_table(ws, r + 2, ("Input completeness controls", "Amount"), controls,
                                bold_prefixes=("Bank: difference", "GL: difference"))

    proof_rows = [("Ending balance per bank statement", dollars(proof["bank_close"]), MONEY)]
    proof_rows += [(f"  {lbl}", dollars(v), MONEY) for lbl, v in proof["bank_lines"]]
    proof_rows += [("Adjusted bank balance", dollars(proof["adj_bank"]), MONEY),
                   ("Ending balance per general ledger", dollars(proof["gl_close"]), MONEY)]
    proof_rows += [(f"  {lbl}", dollars(v), MONEY) for lbl, v in proof["gl_lines"]]
    proof_rows += [("Adjusted GL balance", dollars(proof["adj_gl"]), MONEY),
                   ("Unreconciled difference", dollars(proof["difference"]), MONEY)]
    last = write_label_value_table(ws, r + 2, ("Reconciliation proof", "Amount"), proof_rows,
                                   bold_prefixes=("Adjusted", "Unreconciled", "Ending"))
    status = "TIES" if proof["difference"] == 0 else "DOES NOT TIE — investigate before close"
    ws.cell(row=last + 1, column=1, value=f"Status: {status}").font = Font(
        bold=True, color="375623" if proof["difference"] == 0 else "C00000")
    autosize(ws, [56, 20])

    # ---------------- Exceptions tab
    ws = wb.create_sheet("Exceptions")
    ws.sheet_properties.tabColor = "C00000"
    ws["A1"] = "Exception Report — ranked by dollar exposure"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = "Each break carries a category, a probable cause and the action that clears it."
    ws["A2"].font = SUB_FONT
    headers = ["Rank", "Side", "Date", "Description", "Doc/Ref", "Amount",
               "Category", "Probable cause", "Suggested action", "Source row"]
    rows = [[e["Rank"], e["Side"], e["Date"], e["Description"], e["Doc/Ref"], e["Amount"],
             e["Category"], e["Probable cause"], e["Suggested action"], e["Source row"]]
            for e in exceptions]
    write_table(ws, 4, headers, rows, money_cols={6}, date_cols={3})
    for i, e in enumerate(exceptions, 5):
        fill = CATEGORY_FILLS.get(e["Category"])
        if fill:
            for c in range(1, len(headers) + 1):
                ws.cell(row=i, column=c).fill = fill
    autosize(ws, [6, 10, 12, 42, 13, 14, 18, 50, 52, 10])

    # ---------------- Matched tabs
    def matched_tab(name, color, mlist, note):
        ws = wb.create_sheet(name)
        ws.sheet_properties.tabColor = color
        ws["A1"] = note
        ws["A1"].font = SUB_FONT
        headers = ["Bank date", "Bank description", "Bank amount",
                   "GL date", "GL memo", "GL doc", "GL amount",
                   "Method", "Days Δ", "Amount Δ", "Match score"]
        rows = []
        for m in sorted(mlist, key=lambda x: (bank.at[x["bank_idx"], "Date"], x["bank_idx"])):
            b, g = bank.loc[m["bank_idx"]], gl.loc[m["gl_idx"]]
            rows.append([b["Date"], b["Description"], b["Amount"],
                         g["Date"], g["Memo"], g["DocNo"], g["Amount"],
                         m["method"], m["date_delta"], dollars(m["cents_delta"]),
                         m["score"] if m["score"] is not None else "—"])
        write_table(ws, 3, headers, rows, money_cols={3, 7, 10}, date_cols={1, 4})
        autosize(ws, [12, 38, 14, 12, 42, 12, 14, 13, 8, 11, 12])

    matched_tab("Matched — Exact", "375623", by_pass["Exact"],
                "Pass 1: identical signed amount and identical date.")
    matched_tab("Matched — Timing", "548235", by_pass["Timing"],
                f"Pass 2: identical amount, cleared later — checks by check number; "
                f"others within ±{TIMING_WINDOW_DAYS} days, assigned to minimise total lag.")
    matched_tab("Matched — Tolerance", "70AD47", by_pass["Tolerance"],
                f"Pass 3: amount within ${TOLERANCE_DOLLARS}, ±{TOLERANCE_WINDOW_DAYS} days, "
                f"description similarity ≥ {FUZZY_THRESHOLD}. Amount Δ = residual to write off or adjust.")

    # ---------------- Open Items (aged) — ledger mode only
    if ledger:
        ws = wb.create_sheet("Open Items (aged)")
        ws.sheet_properties.tabColor = "BF8F00"
        ws["A1"] = f"Open reconciling items at {period_label} close"
        ws["A1"].font = TITLE_FONT
        ws["A2"] = ("Every item still open, whichever period raised it, aged from its own "
                    "transaction date. Flagged items need action, not just carrying.")
        ws["A2"].font = SUB_FONT

        summary = [(b["bucket"] + " days", b["count"], dollars(b["cents"]), b["flagged"])
                   for b in ledger["aging_summary"]]
        r = write_table(ws, 4, ["Age bucket", "Items", "Exposure", "Flagged"], summary,
                        money_cols={3}, freeze=False)

        headers = ["Origin period", "Date", "Description", "Doc/Ref", "Amount", "Category",
                   "Age (days)", "Bucket", "Flag", "Escalated"]
        rows = [[a["origin_period_id"], pd.Timestamp(a["origin_date"]), a["description"],
                 a["doc_ref"], dollars(a["amount_cents"]), a["category"], a["age_days"],
                 a["bucket"], a["flag"], "yes" if a["escalated"] else ""]
                for a in ledger["aging"]]
        last = write_table(ws, r + 2, headers, rows, money_cols={5}, date_cols={2})
        for i, a in enumerate(ledger["aging"], r + 3):
            fill = FLAG_FILLS.get(a["flag"])
            if fill:
                for c in range(1, len(headers) + 1):
                    ws.cell(row=i, column=c).fill = fill
        autosize(ws, [14, 12, 44, 14, 14, 18, 11, 10, 11, 10])

        # ---------------- Rollforward
        ws = wb.create_sheet("Rollforward")
        ws.sheet_properties.tabColor = "2F5597"
        ws["A1"] = "Reconciling items — rollforward"
        ws["A1"].font = TITLE_FONT
        ws["A2"] = ("open at start + raised − cleared − written off = open at end. "
                    "One period's closing items are the next one's opening items.")
        ws["A2"].font = SUB_FONT

        headers = ["Period", "Open at start", "Raised", "Cleared", "Written off",
                   "Open at end", "Exposure at end", "Check"]
        rows = [[r_["period_id"], r_["opening_count"], r_["raised_count"], r_["cleared_count"],
                 r_["written_off_count"], r_["closing_count"], dollars(r_["closing_cents"]),
                 r_["status"]] for r_ in ledger["rollforward"]]
        last = write_table(ws, 4, headers, rows, money_cols={7})
        for i, r_ in enumerate(ledger["rollforward"], 5):
            ws.cell(row=i, column=8).font = Font(
                bold=True, color="375623" if r_["status"] == "BALANCED" else "C00000")
        for j, line in enumerate(ledger["breaks"]):
            ws.cell(row=last + 2 + j, column=1, value=f"BREAK: {line}").font = Font(color="C00000")
        if not ledger["breaks"]:
            ws.cell(row=last + 2, column=1,
                    value="No breaks: every period balances and rolls into the next.").font = \
                Font(bold=True, color="375623")
        autosize(ws, [12, 14, 9, 9, 12, 13, 18, 12])

    # ---------------- Source tabs
    def source_tab(name, df, cols, money_col):
        ws = wb.create_sheet(name)
        ws.sheet_properties.tabColor = "808080"
        write_table(ws, 1, cols, df[cols].values.tolist(), money_cols={money_col}, date_cols={1})
        autosize(ws, [12] + [34] * (len(cols) - 2) + [14])

    source_tab("Bank Statement (source)", bank.sort_values(["Date", "src_row"]),
               ["Date", "Description", "Reference", "Amount"], 4)
    source_tab("GL Extract (source)", gl.sort_values(["Date", "src_row"]),
               ["Date", "Account", "Memo", "DocNo", "Amount"], 5)

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path


# ------------------------------------------------------------------ main
def reconcile(bank_path=None, gl_path=None, balances_path=None, conn=None,
              period_id=None, carried=None):
    """Everything except the report — the entry point tests will call.

    Two sources, same result: three CSVs, or a period already in the ledger.
    A parity test pins that both produce identical frames.

    `carried` is the output of carryforward.resolve(): rows it consumed are
    withheld from matching (they belong to an earlier period's item), and
    items it could not clear are carried into the exceptions and the proof."""
    if conn is not None and period_id is not None:
        import db
        bank, gl, balances = db.period_frames(conn, period_id)
    else:
        bank, gl = load(bank_path, gl_path)
        balances = load_balances(balances_path)
    # Rows consumed by pass 0 belong to an earlier period's item, so they are
    # withheld from matching — but they are still rows of THIS file, so the
    # completeness controls below are computed on the full frames.
    bank_m, gl_m = bank, gl
    if carried:
        bank_m = bank[~bank["bank_txn_id"].isin(carried["consumed_bank"])].copy()
        gl_m = gl[~gl["gl_entry_id"].isin(carried["consumed_gl"])].copy()

    matches, bank_open, gl_open = run_matching(bank_m, gl_m)
    exceptions = classify_exceptions(bank_m, gl_m, bank_open, gl_open, matches,
                                     balances["period_end"])
    residual = sum(m["cents_delta"] for m in matches)
    if residual:
        exceptions.append(residual_item(residual, balances["period_end"]))
    if carried:
        exceptions = merge_carried(exceptions, carried["still_open"], balances["period_end"])
    else:
        exceptions.sort(key=lambda r: (-abs(r["cents"]), r["Date"], r["Side"], r["Source row"]))
        for rank, r in enumerate(exceptions, 1):
            r["Rank"] = rank
    proof = build_proof(bank, gl, matches, exceptions, balances)
    return bank, gl, balances, matches, exceptions, proof


def report_path(period_end, out=None, root=None):
    """Name the report after the period it reconciles, so months cannot
    overwrite each other."""
    if out:
        return Path(out)
    root = root or Path(__file__).resolve().parent.parent
    return root / "output" / f"reconciliation_report_{period_end.strftime('%b%Y').lower()}.xlsx"


def print_summary(bank, gl, matches, exceptions, proof, n_carried=None):
    n = defaultdict(int)
    for m in matches:
        n[m["pass"]] += 1
    n_check = sum(1 for m in matches if m["method"] == "check no.")
    n_bank = sum(1 for e in exceptions if e["Side"] == "Bank only")
    n_gl = sum(1 for e in exceptions if e["Side"] == "GL only")

    print(f"Bank rows: {len(bank)}   GL rows: {len(gl)}")
    if n_carried is not None:
        print(f"Pass 0 carryfwd:  {n_carried:>3}")
    print(f"Pass 1 exact:     {n['Exact']:>3}")
    print(f"Pass 2 timing:    {n['Timing']:>3}  "
          f"(check no. {n_check}, date window {n['Timing'] - n_check})")
    print(f"Pass 3 tolerance: {n['Tolerance']:>3}")
    print(f"Total reconciled: {len(matches):>3}")
    print(f"Open items:       {len(exceptions):>3}  (bank-only {n_bank}, GL-only {n_gl})")
    for side, c in proof["completeness"].items():
        print(f"Completeness {side:<4}  {dollars(c['diff']):>12,.2f}  "
              f"{'OK' if c['diff'] == 0 else 'BREAK'}")
    print(f"Unreconciled difference: {dollars(proof['difference']):,.2f}  "
          f"{'TIES' if proof['difference'] == 0 else 'DOES NOT TIE'}")


def run_from_ledger(args, root):
    """Ledger mode: load the period if needed, clear what earlier periods left
    open, reconcile the rest, and save the run back."""
    import carryforward
    import controls
    import db
    import store

    db.init_db(args.db)                                  # no-op if it already exists
    conn = db.connect(args.db)
    try:
        loaded = db.load_period(conn, args.period, args.data_dir, replace=args.replace)
        print(f"Period {args.period}: {loaded['action']} "
              f"(bank {loaded['bank']} rows, GL {loaded['gl']} rows)")

        carried = carryforward.resolve(conn, args.period)
        if carried["resolutions"] or carried["still_open"]:
            print(f"Pass 0 carry-forward: cleared {len(carried['resolutions'])} prior items, "
                  f"{len(carried['still_open'])} still open")

        bank, gl, balances, matches, exceptions, proof = reconcile(
            conn=conn, period_id=args.period, carried=carried)
        print_summary(bank, gl, matches, exceptions, proof, len(carried["resolutions"]))

        run_id = store.save_run(conn, args.period, bank, gl, matches, exceptions, proof,
                                carried=carried)
        n_new = conn.execute("SELECT COUNT(*) FROM reconciling_item WHERE origin_run_id = ?",
                             (run_id,)).fetchone()[0]
        print(f"Saved run {run_id} to the ledger: {len(matches)} matches, "
              f"{n_new} items raised, {len(carried['resolutions'])} cleared")

        # controls run after the items are stored: they read the ledger, not the run
        escalated = controls.apply_escalations(conn, args.period)
        ledger = {"aging": controls.aging(conn, args.period),
                  "aging_summary": controls.aging_summary(conn, args.period),
                  "rollforward": controls.rollforward(conn),
                  "breaks": controls.rollforward_breaks(conn)}
        controls.print_aging(conn, args.period)
        if escalated:
            print(f"  {len(escalated)} newly escalated")
        print()
        controls.print_rollforward(conn)

        out = build_report(bank, gl, matches, exceptions, proof, balances["period_end"],
                           report_path(balances["period_end"], args.out, root),
                           ledger=ledger)
    finally:
        conn.close()
    return out


def run_from_files(args, root):
    """File mode: three CSVs in, one report out. No database involved."""
    missing = [f"--{n}" for n in ("bank", "gl", "balances") if not getattr(args, n)]
    if missing:
        raise SystemExit(f"file mode needs {', '.join(missing)} "
                         f"(or use --period to reconcile from the ledger)")
    bank, gl, balances, matches, exceptions, proof = reconcile(
        args.bank, args.gl, args.balances)
    out = build_report(bank, gl, matches, exceptions, proof, balances["period_end"],
                       report_path(balances["period_end"], args.out, root))
    print_summary(bank, gl, matches, exceptions, proof)
    return out


def main():
    root = Path(__file__).resolve().parent.parent
    ap = argparse.ArgumentParser(
        description="Bank-to-GL reconciliation: carry-forward + three matching passes",
        epilog="examples:  reconcile.py --period 2026-07   |   "
               "reconcile.py --bank b.csv --gl g.csv --balances bal.csv")
    ap.add_argument("--period", help="reconcile this period from the ledger (e.g. 2026-07) "
                                     "and save the run back into it")
    ap.add_argument("--bank", help="bank statement CSV (file mode)")
    ap.add_argument("--gl", help="GL cash extract CSV (file mode)")
    ap.add_argument("--balances", help="stated opening/closing balances CSV (file mode)")
    ap.add_argument("--out", help="report path "
                                  "(default: output/reconciliation_report_<period>.xlsx)")
    ap.add_argument("--db", default=str(root / "ledger.db"))
    ap.add_argument("--data-dir", default=str(root / "data"),
                    help="where this period's CSVs live (ledger mode)")
    ap.add_argument("--replace", action="store_true",
                    help="reload this period's CSVs even if the stored rows differ")
    args = ap.parse_args()

    out = run_from_ledger(args, root) if args.period else run_from_files(args, root)
    print(f"Report: {out}")


if __name__ == "__main__":
    main()